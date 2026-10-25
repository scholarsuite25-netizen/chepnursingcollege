import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_future_nursing_ai_standard_docx():
    doc = Document()
    
    # Page setup - Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)
        
    NAVY = RGBColor(16, 42, 77)          # #102A4D
    EMERALD = RGBColor(16, 122, 85)      # #107A55
    PURPLE = RGBColor(107, 33, 168)      # #6B21A8
    DARK_GRAY = RGBColor(45, 55, 72)     # #2D3748
    MUTED_GRAY = RGBColor(100, 116, 139) # #64748B
    
    # Document Header Eyebrow
    eyebrow = doc.add_paragraph()
    eyebrow.paragraph_format.space_before = Pt(0)
    eyebrow.paragraph_format.space_after = Pt(2)
    run_eye = eyebrow.add_run("GLOBAL CLINICAL ACCREDITATION STANDARD • AI-NEXS 2026.1")
    run_eye.font.name = 'Arial'
    run_eye.font.size = Pt(10)
    run_eye.font.bold = True
    run_eye.font.color.rgb = PURPLE
    
    # Title
    main_title = doc.add_paragraph()
    main_title.paragraph_format.space_before = Pt(2)
    main_title.paragraph_format.space_after = Pt(4)
    run_t = main_title.add_run("THE FUTURE OF NURSING CLINICAL EXAMINATIONS & ARTIFICIAL INTELLIGENCE")
    run_t.font.name = 'Arial'
    run_t.font.size = Pt(20)
    run_t.font.bold = True
    run_t.font.color.rgb = NAVY
    
    # Subtitle
    sub = doc.add_paragraph()
    sub.paragraph_format.space_before = Pt(0)
    sub.paragraph_format.space_after = Pt(12)
    run_sub = sub.add_run("Comprehensive Regulatory Framework, Deficit Remediation, Multi-Modal Modalities & Algorithmic Stewardship for Testing the Nurse of the Future")
    run_sub.font.name = 'Arial'
    run_sub.font.size = Pt(11)
    run_sub.font.italic = True
    run_sub.font.color.rgb = MUTED_GRAY

    # Metadata Banner Table
    banner_table = doc.add_table(rows=1, cols=1)
    banner_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    banner_cell = banner_table.cell(0, 0)
    set_cell_background(banner_cell, "F3E8FF")  # Light Purple
    set_cell_margins(banner_cell, 160, 160, 220, 220)
    
    bp = banner_cell.paragraphs[0]
    bp.paragraph_format.space_before = Pt(0)
    bp.paragraph_format.space_after = Pt(0)
    r1 = bp.add_run("CO-ACCREDITATION & REGULATORY HARMONIZATION: ")
    r1.font.name = 'Arial'
    r1.font.size = Pt(9.5)
    r1.font.bold = True
    r1.font.color.rgb = PURPLE
    r2 = bp.add_run("Nursing and Midwifery Council of Nigeria (NMCN) Act Cap N143 LFN 2004 • NCSBN Clinical Judgment Model (Next-Gen NCLEX) • International Council of Nurses (ICN) • WHO Global Strategic Directions for Nursing and Midwifery 2021-2025 • AMEE Guide No. 81.")
    r2.font.name = 'Arial'
    r2.font.size = Pt(9.5)
    r2.font.color.rgb = DARK_GRAY

    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(4)
    p_spacer.paragraph_format.space_after = Pt(4)

    # Helper function for headings
    def add_custom_heading(text, level=1):
        h = doc.add_paragraph()
        run = h.add_run(text)
        run.font.name = 'Arial'
        run.font.bold = True
        if level == 1:
            h.paragraph_format.space_before = Pt(16)
            h.paragraph_format.space_after = Pt(6)
            run.font.size = Pt(14)
            run.font.color.rgb = NAVY
        elif level == 2:
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(4)
            run.font.size = Pt(12)
            run.font.color.rgb = EMERALD
        else:
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(2)
            run.font.size = Pt(10.5)
            run.font.color.rgb = PURPLE
        return h

    def add_body(text, bold_prefix=None, italic=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            rb = p.add_run(bold_prefix)
            rb.font.name = 'Arial'
            rb.font.size = Pt(10)
            rb.font.bold = True
            rb.font.color.rgb = DARK_GRAY
        rt = p.add_run(text)
        rt.font.name = 'Arial'
        rt.font.size = Pt(10)
        rt.font.italic = italic
        rt.font.color.rgb = DARK_GRAY
        return p

    # 1. Executive Summary
    add_custom_heading("1. EXECUTIVE CHARTER & THE DILEMMA IN MODERN CLINICAL TESTING", 1)
    add_body(
        "For nearly five decades, clinical nursing licensure and institutional evaluations worldwide have relied almost exclusively on static Multiple-Choice Questions (MCQs), paper-and-pencil essays, and physical Objective Structured Clinical Examinations (OSCE). While these formats successfully tested rote recall and ritualized motor skills, international educational bodies (WHO, ICN, NCSBN, and NMCN) have declared them profoundly deficient for preparing nurses to practice safely in increasingly complex, technology-driven clinical ecosystems.",
        "The Paradigm Crisis: "
    )
    add_body(
        "Modern healthcare environments are characterized by rapid clinical deterioration, polypharmacy, multi-morbidity, electronic health records (EHR), and the ubiquity of Artificial Intelligence Clinical Decision Support (CDS) algorithms. In this setting, ritualized memorization without dynamic clinical judgment leads to failure to rescue, medication calculation catastrophes, and missed nursing care. This standard formalizes the next-generation paradigm for evaluating the 'Nurse of the Future'."
    )

    # 2. The 7 Gaps in Contemporary Nursing Exams
    add_custom_heading("2. CRITICAL GAPS CURRENTLY ENVISAGED & DECLARED MISSING IN NURSING ASSESSMENTS", 1)
    
    gaps = [
        ("The 'Checklist Robot' Syndrome: ", "Current physical OSCE examinations evaluate procedural compliance (e.g. wash hands for 20 seconds, pick up tray) rather than adaptive diagnostic clinical reasoning. Candidates learn to memorize step sequences without understanding underlying pathophysiology."),
        ("Static, Frozen Vital Signs: ", "In traditional examinations, a patient's vitals remain unchanged regardless of candidate intervention. Real patients exhibit immediate physiological feedback—failing to elevate oxygen saturation leads to hypoxic bradycardia; excessive fluid boluses induce flash pulmonary edema."),
        ("Subjectivity and Inter-Rater Reliability Variance: ", "Examiner fatigue, cognitive bias, and geographic variability create unacceptable discrepancies across examination centers, undermining statutory credential integrity."),
        ("Absence of Algorithmic Stewardship Assessment: ", "Bedside nurses now routinely interact with AI clinical decision support systems in EHRs. No traditional exam tests whether a nurse can detect an AI hallucination, biased triage score, or inappropriate drug dosing recommendation."),
        ("Fragmented Examination Modalities in Isolated Silos: ", "Dosage math, diagnostic interpretation, labor monitoring, and communication are tested in disparate rooms at different times, rather than evaluating unified shift management."),
        ("Zero Longitudinal Shift Unfolding: ", "Traditional OSCE stations last only 5 to 10 minutes. Real clinical practice requires managing evolving conditions across an 8-to-12 hour hospital shift."),
        ("Slow, Non-Actionable Remediation Feedback: ", "Students wait weeks or months for exam results, receiving a single percentage score with zero diagnostic feedback on specific clinical reasoning deficits.")
    ]
    for b, t in gaps:
        add_body(t, b)

    # 3. The 7 Pillars Surpassing Global Standards
    add_custom_heading("3. THE 7 ARCHITECTURAL PILLARS THAT SURPASS GLOBAL STANDARDS", 1)
    
    # Table of Pillars
    t_pillars = doc.add_table(rows=1, cols=3)
    t_pillars.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t_pillars.rows[0].cells
    hdr[0].text = "Pillar / Modality"
    hdr[1].text = "Technological Innovation"
    hdr[2].text = "Clinical Outcome Evaluated"
    for c in hdr:
        set_cell_background(c, "102A4D")
        for p in c.paragraphs:
            p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
            p.runs[0].font.bold = True
            p.runs[0].font.size = Pt(9.5)

    pillar_data = [
        ("Pillar 1: Dynamic Hemodynamic Engine", "Mathematical state-machine linking vitals to interventions in real time.", "Immediate rescue of deteriorating patients; prevention of clinical shock."),
        ("Pillar 2: Multimodal Conversational Patient", "Speech-to-text / interactive dialogue simulating complex patient interviews.", "Therapeutic communication, active listening, and SBAR physician handover."),
        ("Pillar 3: Algorithmic Stewardship & AI Oversight", "Simulated EMR with intentional AI diagnostic hallucinations and drug errors.", "Testing nurse vigilance to audit, override, and correct flawed automated CDS orders."),
        ("Pillar 4: Item Response Theory (IRT & CAT)", "3-Parameter Logistic adaptive engine dynamically tailoring item difficulty.", "Ultra-precise ability estimation (Theta); elimination of question leaks."),
        ("Pillar 5: Cryptographic Trust Architecture", "W3C Verifiable Credentials with SHA-256 tokens and verifiable QR codes.", "Zero credential forgery; instant employer and NMCN verification."),
        ("Pillar 6: Cognitive Prescription Remediation", "Automated diagnostic feedback prescribing exact textbook chapters & simulators.", "Elimination of repeat failures; closed-loop mastery learning."),
        ("Pillar 7: Unified Multi-Modal Studio", "CAOSCE, Partograph, OSPE, Dosage Math, NGN EHR, Care Plan, START Triage.", "Holistic 360-degree assessment of clinical readiness.")
    ]

    for p, tech, out in pillar_data:
        row = t_pillars.add_row().cells
        set_cell_background(row[0], "F8FAFC")
        set_cell_margins(row[0], 80, 80, 100, 100)
        set_cell_margins(row[1], 80, 80, 100, 100)
        set_cell_margins(row[2], 80, 80, 100, 100)
        row[0].paragraphs[0].add_run(p).font.bold = True
        row[0].paragraphs[0].runs[0].font.size = Pt(9)
        row[1].paragraphs[0].add_run(tech).font.size = Pt(9)
        row[2].paragraphs[0].add_run(out).font.size = Pt(9)

    p_spacer2 = doc.add_paragraph()
    p_spacer2.paragraph_format.space_before = Pt(6)
    p_spacer2.paragraph_format.space_after = Pt(6)

    # 4. Multi-Modal Assessment Catalog
    add_custom_heading("4. THE COMPLETE MULTI-MODAL CLINICAL EXAMINATION SUITE", 1)
    
    modalities = [
        ("1. 6-Station Autonomous CAOSCE Circuit: ", "Six sequential clinical domains (Cardiopulmonary Resuscitation, Active Management of 3rd Stage Labor, Aseptic Wound Care, Pediatric Asthma & Respiratory Distress, Post-Operative Blood Transfusion Reaction, and Cervical Spine Trauma Management). Each station features a 10-minute practical checklist followed by 10 linked MCQs (60 MCQs total) under automated bell chimes."),
        ("2. Interactive WHO Modified Partograph Simulator: ", "Vector plotting of cervical dilation and fetal head descent against statutory Alert and Action lines. Evaluates maternal pulse, blood pressure, liquor status, and contraction duration, triggering real-time transfer and cesarean section decision timing."),
        ("3. OSPE Diagnostic Data Studio: ", "Objective Structured Practical Examination evaluating 12-lead ECG rhythm interpretation, Arterial Blood Gas (ABG) compensation, Urinalysis dipstick colorimetry, Chest Radiograph consolidation, Surgical Instrumentation identification, and Electrolyte/Renal lab panels."),
        ("4. Zero-Tolerance Dosage & IV Fluid Calculation Board: ", "Mandatory 100% precision testing for IV gravity drip rates (gtt/min), pediatric weight-based suspensions (mg/kg), Parkland Burn Resuscitation fluid protocols, and continuous Heparin anticoagulation titration."),
        ("5. Unfolding Next-Gen NCLEX (NGN) EHR Case Study Simulator: ", "Longitudinal simulated Electronic Health Record containing Nurses' Notes, Vital Signs Flowsheet, Laboratory Panels, and Provider Orders. Evaluates candidates using Highlight Cue Recognition, Matrix Multiple-Response grids, and 6-part Bow-Tie clinical reasoning diagrams."),
        ("6. Nursing Care Plan (NCP) 5-Step Clinical Reasoning Matrix: ", "Formulation of complete 5-step nursing process (Assessment sorting, NANDA-I prioritized diagnostic formulation, NOC SMART expected outcomes, NIC evidence-based interventions, and Hour-3 clinical evaluation scoring)."),
        ("7. START Disaster Mass-Casualty Triage Simulator: ", "60-second rapid triage of mass-casualty disaster incidents (building collapses, tanker explosions, industrial fires), categorizing casualties into RED (Immediate P1), YELLOW (Delayed P2), GREEN (Minor P3), and BLACK (Deceased P0)."),
        ("8. Conversational AI Clinical Preceptor & SBAR Handover Simulator: ", "Interactive dialogue encounters assessing patient history taking, pain scale assessment (PQRST), allergy verification, structured SBAR communication with physicians, and Algorithmic Stewardship error override.")
    ]
    for b, t in modalities:
        add_body(t, b)

    # 5. Regulatory Gating & Scoring Architecture
    add_custom_heading("5. REGULATORY SCORING & PASS/FAIL CRITERIA", 1)
    add_body(
        "In accordance with NMCN statutory regulations, a candidate must achieve a minimum composite score of 50.0% to pass any examination. A score >= 75.0% constitutes an institutional Distinction. However, commission or omission of a Critical Patient Safety Step (such as failing to exclude a second twin prior to Oxytocin administration, or performing a head-tilt in suspected cervical trauma) constitutes an AUTOMATIC ZERO for the station, regardless of MCQ performance.",
        "Statutory Thresholds: "
    )

    # Save to workspace root and artifacts directory
    workspace_docx = "c:/Users/Registrar/Music/stitch_nursing_college_next.js_migration/GLOBAL_FUTURE_NURSING_EXAMINATION_AI_STANDARD.docx"
    artifact_docx = "C:/Users/Registrar/.gemini/antigravity-ide/brain/ae1fec37-b909-493c-acd3-c76fde7b4686/GLOBAL_FUTURE_NURSING_EXAMINATION_AI_STANDARD.docx"
    
    doc.save(workspace_docx)
    doc.save(artifact_docx)
    print(f"Generated DOCX at {workspace_docx} and {artifact_docx}")

if __name__ == "__main__":
    create_future_nursing_ai_standard_docx()
